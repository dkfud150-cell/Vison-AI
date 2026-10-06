"""
문서 출력 — 위험성평가표(엑셀), 시정지시서 · 작업허가서 점검 · 산업재해조사표(워드)

모든 문서 끝에 '누가 무엇을 채웠는지'(코드 / AI 초안 / 사람)와 보존 근거를 적는다.
"""
from __future__ import annotations

import time
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import zone_name

GRADE_FILL = {"낮음": "DCEFE2", "보통": "FFF1C2", "높음": "FAD7B5", "매우 높음": "F4C2BD"}
THIN = Side(style="thin", color="9AA0A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
HEAD_FILL = PatternFill("solid", fgColor="E8EBEF")
FONT = "맑은 고딕"


def _safe_save(save_fn, path: Path) -> Path:
    """대상 파일이 다른 프로그램(엑셀·워드, OneDrive 동기화 등)에 의해 잠겨 있으면
    타임스탬프를 붙인 이름으로 대신 저장한다."""
    path = Path(path)
    try:
        save_fn(path)
        return path
    except PermissionError:
        alt = path.with_stem(f"{path.stem}_{time.strftime('%H%M%S')}")
        save_fn(alt)
        return alt


def _measures_text(measures):
    return "\n".join(f"[{m['category']}] {m['text']}" for m in measures)


def assessment_to_excel(drafts, path, *, rulesets, scales, site="산세 공장",
                        record: dict | None = None) -> Path:
    """
    위험성평가표 엑셀. 시트 3장
      1) 위험성평가표 — 단위작업 · 유해위험요인 · 현재 조치 · 위험도 · 근거 이벤트 · 감소대책 · 개선 후(실측)
      2) 근거 이벤트   — 행마다 어떤 이벤트가 근거였는지 (클립 경로 포함)
      3) 척도 기준     — 가능성 5 × 중대성 4 환산표
    record: assessment.get_assessment() 결과. 있으면 서명 상태를 적는다.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "위험성평가표"

    period = f"{drafts[0]['window_from'][:10]} ~ {drafts[0]['window_to'][:10]}" if drafts else "-"
    ws["A1"] = "위험성평가표 (빈도·강도법, 가능성 5 × 중대성 4)"
    ws["A1"].font = Font(name=FONT, size=15, bold=True)
    kind = (record or {}).get("kind", "수시 · 정기 공용")
    ws["A2"] = f"사업장: {site}    평가 근거 기간: {period}    구분: {kind}"
    if not record:
        state = "초안 — 관리자 서명 · 근로자 대표 확인 전"
    elif record["state"] == "confirmed":
        state = f"확정 {record['ra_id']} ({record['confirmed_at']})"
    else:
        state = f"{record['state_label']} {record['ra_id']}"
    ws["A3"] = "상태: " + state
    for c in ("A2", "A3"):
        ws[c].font = Font(name=FONT, size=10, color="555555")

    headers = ["No", "단위작업\n(작업 모드)", "구역", "유해·위험요인", "발생 원인", "현재 안전조치",
               "가능성", "중대성", "위험도", "등급", "근거 이벤트\n(최근 90일)", "감소대책",
               "규칙 변경", "개선 후 위험도\n(적용 후 실측)", "법적 근거", "수시평가 / 확인 필요"]
    widths = [5, 18, 12, 26, 32, 26, 7, 7, 7, 9, 18, 42, 9, 18, 30, 28]
    hr = 5
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=hr, column=i, value=h)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
        cell.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hr].height = 34

    for n, d in enumerate(drafts, 1):
        r = hr + n
        basis = f"기본 {d['base_count']}건\n(최근 14일 {d['recent_count']}건)"
        if d["adj_count"]:
            basis += f"\n강화 기간 {d['adj_count']}건 별도"
        if d.get("top_alert_count"):
            basis += f"\n최고 경보 {d['top_alert_count']}건"
        if d.get("accident_ids"):
            basis += f"\n사고 {', '.join(d['accident_ids'])}"
        flags = []
        if d.get("adhoc_reasons"):
            flags.append("수시평가 제안: " + ", ".join(d["adhoc_reasons"]))
        if d.get("check_needed"):
            flags.append("확인: " + d["check_needed"])
        flags += [f"주의: {w}" for w in d.get("warnings", [])]
        values = [
            n, f"{d.get('unit_work', '')}\n({d['mode']})",
            f"{d['zone']}\n{zone_name(rulesets, d['zone'])}",
            d.get("hazard", ""), d.get("cause", ""),
            "\n".join(d.get("current_measures", [])),
            d["likelihood"], d["severity"], d["score"], d["grade"], basis,
            _measures_text(d.get("measures", [])),
            ", ".join(f"{i} (반려)" if i in (record or {}).get("rule_rejected", {}) else i
                      for i in sorted({m["rule_id"] for m in d.get("measures", []) if m.get("rule_id")})) or "-",
            (d.get("after") or {}).get("text") or "적용 후 실측\n(대책 적용일 전후 위반 건수)",
            "\n".join(d.get("legal_refs", [])), "\n".join(flags),
        ]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.border = BOX
            cell.alignment = CENTER if c in (1, 7, 8, 9, 10, 13) else WRAP
        ws.cell(row=r, column=10).fill = PatternFill("solid", fgColor=GRADE_FILL.get(d["grade"], "FFFFFF"))
        ws.cell(row=r, column=9).font = Font(name=FONT, size=11, bold=True)
        ws.cell(row=r, column=14).font = Font(name=FONT, size=9, color="666666")
        if d.get("adhoc_reasons"):
            ws.cell(row=r, column=16).font = Font(name=FONT, size=10, bold=True, color="9D2416")
        lines = max(4, len(_measures_text(d.get("measures", []))) // 22 + 2)
        ws.row_dimensions[r].height = min(15 * lines, 170)

    # 확인란 — 산업안전보건법 제36조제2항 · 위험성평가 지침 제6조 근로자 참여
    rec = record or {}
    sr = hr + len(drafts) + 2
    ws.cell(row=sr, column=2, value="확인").font = Font(name=FONT, bold=True)
    ws.cell(row=sr + 1, column=2, value="관리자 (검토·서명)")
    ws.cell(row=sr + 1, column=4, value=rec.get("manager") or "(서명)")
    ws.cell(row=sr + 2, column=2, value="근로자 대표 (확인)")
    ws.cell(row=sr + 2, column=4, value=rec.get("worker_rep") or "(서명)")
    ws.cell(row=sr + 4, column=2,
            value="※ 가능성·중대성·위험도는 감지 이벤트로 시스템이 계산한 값이다. 문장과 감소대책은 AI 초안이며, "
                  "관리자 서명과 근로자 대표 확인을 모두 거쳐야 평가서로 성립한다(산업안전보건법 제36조제2항). "
                  "개선 후 위험도는 예측하지 않고, 대책 적용일 전후 같은 기간의 위반 건수로 실측한다.")
    ws.cell(row=sr + 5, column=2,
            value="※ 보존: 3년 (산업안전보건법 시행규칙 제37조제2항). 기산일은 평가를 완료한 날(위험성평가 지침 제14조제2항).")
    for rr in (sr + 4, sr + 5):
        ws.cell(row=rr, column=2).font = Font(name=FONT, size=9, color="666666")
    ws.freeze_panes = ws.cell(row=hr + 1, column=3)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    # 시트 2 — 근거 이벤트
    ws2 = wb.create_sheet("근거 이벤트")
    h2 = ["평가표 No", "구역", "작업 모드", "위반 유형", "구분", "이벤트 번호"]
    for i, h in enumerate(h2, 1):
        c = ws2.cell(row=1, column=i, value=h)
        c.font = Font(name=FONT, bold=True)
        c.fill = HEAD_FILL
        c.border = BOX
    ws2.column_dimensions["A"].width = 10
    for col, w in zip("BCDEF", [10, 22, 22, 12, 60]):
        ws2.column_dimensions[col].width = w
    r = 2
    for n, d in enumerate(drafts, 1):
        for label, ids in (("기본 규칙", d["event_ids"]), ("강화 기간", d["adj_event_ids"])):
            if not ids:
                continue
            for c, v in enumerate([n, d["zone"], d["mode"], d["violation_type"], label, ", ".join(ids)], 1):
                cell = ws2.cell(row=r, column=c, value=v)
                cell.font = Font(name=FONT, size=10)
                cell.border = BOX
                cell.alignment = WRAP
            r += 1
    ws2.cell(row=r + 1, column=1, value="클립은 clips/<이벤트 번호>.mp4 (전후 10초). 무음 기록·오탐 표시 건은 빈도에서 제외.")

    # 시트 3 — 척도 기준
    ws3 = wb.create_sheet("척도 기준")
    ws3["A1"] = (f"가능성 (최근 {scales['window_days']}일 기본 규칙 위반 건수 — 무음 기록·오탐·강화 기간 제외)")
    ws3["A1"].font = Font(name=FONT, bold=True)
    row = 2
    for st in scales["likelihood_steps"]:
        ws3.cell(row=row, column=1, value=st["level"])
        ws3.cell(row=row, column=2, value=st["label"])
        row += 1
    fl = scales.get("likelihood_floor", {})
    ws3.cell(row=row, column=2, value=f"하한: 최고 경보 1건 → 최소 {fl.get('top_alert', 4)}, "
                                      f"사고와 이어진 조합 → {fl.get('accident', 5)}")
    row += 1
    row += 1
    ws3.cell(row=row, column=1, value="중대성 (위반 유형별)").font = Font(name=FONT, bold=True)
    row += 1
    for vt, sv in scales["severity_by_violation"].items():
        if vt.startswith("_"):              # _설명 같은 주석 줄
            continue
        ws3.cell(row=row, column=1, value=sv)
        ws3.cell(row=row, column=2, value=vt)
        ws3.cell(row=row, column=3, value=scales["severity_labels"][str(sv)])
        row += 1
    for mode, fv in scales.get("severity_floor_by_mode", {}).items():
        if isinstance(fv, int):
            ws3.cell(row=row, column=2, value=f"모드 하한: {mode} → 최소 {fv}")
            row += 1
    row += 1
    ws3.cell(row=row, column=1, value="등급 (가능성 × 중대성)").font = Font(name=FONT, bold=True)
    row += 1
    lo = 1
    for g in scales["grades"]:
        ws3.cell(row=row, column=1, value=f"{lo}~{g['max_score']}")
        ws3.cell(row=row, column=2, value=g["label"])
        ws3.cell(row=row, column=3, value=g["action"])
        ws3.cell(row=row, column=2).fill = PatternFill("solid", fgColor=GRADE_FILL[g["label"]])
        lo = g["max_score"] + 1
        row += 1
    for col, w in zip("ABC", [12, 26, 40]):
        ws3.column_dimensions[col].width = w

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    return _safe_save(wb.save, path)


# ---------------------------------------------------------------------- 워드
def _set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _col_widths(table, widths_cm):
    """LibreOffice·워드 모두에서 열 너비가 먹도록 모든 칸에 너비를 지정한다."""
    table.autofit = False
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))       # 1cm = 567 twips
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def _korean_font(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def corrective_to_docx(order: dict, path) -> Path:
    doc = Document()
    _korean_font(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("시 정 지 시 서")
    run.bold = True
    run.font.size = Pt(20)
    p = doc.add_paragraph(f"문서번호 {order['doc_no']}    발행 {order['issued']}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    ev = order["event"]
    doc.add_heading("1. 위반 개요", level=2)
    info = [
        ("대상 구역", f"{ev['zone']} {order['zone_name']}"),
        ("작업 모드", ev["mode"]),
        ("감지 일시", ev["ts"].replace("T", " ")),
        ("감지 유형", ev["violation_type"]),
        ("적용 규칙", order["rule_state"]),
        ("같은 위반 최근 14일", f"{order['same_recent']}건"),
        ("중대성", f"{order['severity']}등급 — {order['severity_label']}"),
        ("근거 기록", f"이벤트 {ev['event_id']} / 클립 {ev.get('clip_path') or '-'}"),
    ]
    tb = doc.add_table(rows=0, cols=2)
    tb.style = "Table Grid"
    for k, v in info:
        cells = tb.add_row().cells
        cells[0].text, cells[1].text = k, v
        _set_cell_bg(cells[0], "E8EBEF")
    _col_widths(tb, [4, 13])

    doc.add_heading("2. 현장 상황 (CCTV 프레임 분석)", level=2)
    if order.get("frame_path"):
        doc.add_picture(order["frame_path"], width=Cm(12))
    doc.add_paragraph(order["scene"] or "-")

    doc.add_heading("3. 위험 기인물·요인", level=2)
    ht = doc.add_table(rows=1, cols=3)
    ht.style = "Table Grid"
    for c, h in zip(ht.rows[0].cells, ["순위", "기인물·요인", "이유"]):
        c.text = h
        _set_cell_bg(c, "E8EBEF")
    for h in order["hazards"]:
        cells = ht.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = str(h.get("rank", "")), h["hazard"], h["reason"]
    _col_widths(ht, [1.5, 6, 9.5])

    doc.add_heading("4. 위반 사항", level=2)
    doc.add_paragraph(order["violation"] or "-")

    doc.add_heading("5. 시정 요구 사항", level=2)
    at = doc.add_table(rows=1, cols=3)
    at.style = "Table Grid"
    for c, h in zip(at.rows[0].cells, ["구분", "조치 내용", "기한"]):
        c.text = h
        _set_cell_bg(c, "E8EBEF")
    for a in order["actions"]:
        cells = at.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = a["category"], a["text"], order["deadline"]
    _col_widths(at, [2.2, 10.3, 4.5])

    doc.add_heading("6. 법적 근거", level=2)
    for ref in order["legal_refs"] or ["-"]:
        doc.add_paragraph(ref, style="List Bullet")

    if order.get("uncertain") or order.get("warnings"):
        doc.add_heading("7. 확인 필요", level=2)
        if order.get("uncertain"):
            doc.add_paragraph(order["uncertain"], style="List Bullet")
        for w in order.get("warnings", []):
            doc.add_paragraph(w, style="List Bullet")

    doc.add_heading("조치 결과 회신", level=2)
    rt = doc.add_table(rows=3, cols=2)
    rt.style = "Table Grid"
    for row, k, h in zip(rt.rows, ["조치 완료일", "조치 내용·사진", "확인자"], [0.9, 3.0, 0.9]):
        row.cells[0].text = k
        row.height = Cm(h)
        _set_cell_bg(row.cells[0], "E8EBEF")
    _col_widths(rt, [4, 13])

    doc.add_paragraph()
    sig = doc.add_paragraph("발행: 안전관리자 ____________ (서명)    수신: 해당 작업 관리감독자 ____________ (서명)")
    sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _note(doc, "※ 현장 상황·위험요인·시정 조치 문장은 AI 초안이다. 조치 기한과 근거 기록은 시스템이 채웠다. "
               "발행 전 안전관리자가 검토한다.")
    _note(doc, "※ 보존: 조치 이행일부터 5년 (중대재해 처벌 등에 관한 법률 시행령 제13조).")
    return _save_docx(doc, path)


def _note(doc, text):
    p = doc.add_paragraph(text)
    for r in p.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def _save_docx(doc, path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    return _safe_save(doc.save, path)


def _new_doc(title: str, sub: str):
    doc = Document()
    _korean_font(doc)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(1.8)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    p = doc.add_paragraph(sub)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return doc


def _kv_table(doc, pairs, widths=(4, 13)):
    tb = doc.add_table(rows=0, cols=2)
    tb.style = "Table Grid"
    for k, v in pairs:
        cells = tb.add_row().cells
        cells[0].text, cells[1].text = k, str(v)
        _set_cell_bg(cells[0], "E8EBEF")
    _col_widths(tb, list(widths))
    return tb


def _grid(doc, headers, rows, widths):
    tb = doc.add_table(rows=1, cols=len(headers))
    tb.style = "Table Grid"
    for c, h in zip(tb.rows[0].cells, headers):
        c.text = h
        _set_cell_bg(c, "E8EBEF")
    for r in rows:
        cells = tb.add_row().cells
        for c, v in zip(cells, r):
            c.text = str(v)
    _col_widths(tb, widths)
    return tb


def ptw_to_docx(result: dict, path) -> Path:
    """작업허가서 점검 결과. 판정은 코드, 지적 문장은 AI 초안(또는 템플릿)."""
    ptw = result["ptw"]
    doc = _new_doc("작업허가서 점검 결과", f"{ptw['ptw_id']}    점검 {result['checked_at'].replace('T', ' ')}")
    doc.add_heading("1. 허가 내용", level=2)
    _kv_table(doc, [("유형", ptw["type"]), ("구역", ptw["zone"]), ("작업", ptw["work"]),
                    ("기간", f"{ptw['start'].replace('T', ' ')} ~ {ptw['end'].replace('T', ' ')}"),
                    ("작업책임자", ptw.get("supervisor") or "(비어 있음)"), ("인원", ptw.get("workers", "-")),
                    ("퍼지", f"{(ptw.get('purge') or {}).get('minutes', '-')}분"),
                    ("화재감시자", ptw.get("fire_watch") or "-")])

    if result["pid_lines"]:
        doc.add_heading("2. 배관 계통도 ↔ 격리 목록 대조", level=2)
        listed = {r.get("line_id"): r for r in result["iso_rows"]}
        rows = []
        for ln in result["pid_lines"]:
            r = listed.get(ln["line_id"])
            yn = lambda v: {"O": "예", "X": "아니오"}.get(str(v or "").strip().upper(), v or "-")
            rows.append([ln["line_id"], ln["name"], ln["fluid"],
                         "있음" if r else "없음 ◀", (r or {}).get("격리방법", "-"),
                         yn((r or {}).get("맹판설치")), yn((r or {}).get("가스측정"))])
        _grid(doc, ["배관", "이름", "유체", "격리 목록", "격리 방법", "맹판", "측정"], rows,
              [1.6, 5.2, 1.8, 2.2, 2.2, 1.6, 1.6])

    doc.add_heading("3. 점검 결과 — " + result["result"], level=2)
    doc.add_paragraph(result.get("summary", ""))
    if result["findings"]:
        _grid(doc, ["항목", "찾은 내용", "지적 사항", "근거"],
              [[f"{f['code']} {f['title']}", f["detail"], f.get("note", ""), "\n".join(f.get("legal_refs", []))]
               for f in result["findings"]], [3.2, 4.2, 6, 3.6])
    doc.add_heading("보완 확인", level=2)
    _kv_table(doc, [("보완 완료 시각", ""), ("보완 내용", ""), ("작업책임자 확인", ""), ("안전관리자 확인", "")])
    doc.add_paragraph()
    _note(doc, f"※ 판정(항목·수치)은 시스템이 규칙셋 isolation 기준으로 계산했다. 지적 문장은 {result.get('text_source', 'AI')} 작성이다.")
    _note(doc, "※ 작업허가서의 법정 보존 기간은 없다. 이 점검 기록은 유해·위험요인 개선 절차의 이행 증빙으로 5년 둔다 "
               "(중대재해 처벌 등에 관한 법률 시행령 제4조제3호 · 제13조).")
    return _save_docx(doc, path)


def accident_to_docx(report: dict, path, profile: dict | None = None) -> Path:
    """산업재해조사표 초안 — 시행규칙 별지 제30호서식의 큰 항목을 따른다. 재해자 정보는 사람이 채운다."""
    acc = report["acc"]
    doc = _new_doc("산업재해조사표 (초안)", f"{acc['id']}    작성 {report['drafted_at'].replace('T', ' ')}    제출 기한 {report['due']}")
    doc.add_heading("1. 사업장 정보", level=2)
    _kv_table(doc, [("사업장명", (profile or {}).get("site_name", "")), ("업종", (profile or {}).get("industry", "")),
                    ("근로자 수", (profile or {}).get("workers", "")), ("사업자등록번호 · 소재지 · 연락처", "(사람이 작성)")])
    doc.add_heading("2. 재해 정보", level=2)
    _kv_table(doc, [("재해자 인적사항", "(사람이 작성 — 이 시스템은 개인 정보를 다루지 않는다)"),
                    ("부상자 수", acc.get("injured", "")), ("휴업 예상 일수", acc.get("lost_days", "")),
                    ("사망 여부", "예" if acc.get("fatal") else "아니오")])
    doc.add_heading("3. 재해 발생 개요", level=2)
    _kv_table(doc, [("발생 일시", report["at"].replace("T", " ")), ("발생 장소", report["zone_label"]),
                    ("적용 중이던 자동 조정", ", ".join(report["active_adjustments"]) or "없음 (기본 규칙)")])
    doc.add_paragraph()
    doc.add_paragraph("재해 발생 경과 (AI 초안)").runs[0].bold = True
    doc.add_paragraph(report["narrative"])
    doc.add_paragraph(f"관제 기록 타임라인 (사고 전 {report['lookback_hours']}시간)").runs[0].bold = True
    _grid(doc, ["시각", "기록", "내용"],
          [[t["ts"].replace("T", " ")[5:16], t["src"], t["what"]] for t in report["timeline"]], [3, 3.5, 10.5])
    doc.add_heading("4. 재해 발생 원인", level=2)
    doc.add_paragraph("(조사자가 작성 — AI 초안은 원인을 단정하지 않는다)")
    doc.add_heading("5. 재발방지 계획 (AI 초안)", level=2)
    _grid(doc, ["구분", "내용"], [[p["category"], p["text"]] for p in report["prevention"]], [3, 14])
    if report.get("check_needed"):
        doc.add_heading("조사자 확인 필요", level=2)
        doc.add_paragraph(report["check_needed"])
    doc.add_paragraph()
    _note(doc, f"※ 발생 일시·장소·타임라인은 관제 기록에서 시스템이 채웠다. 경과·재발방지 문장은 {report.get('text_source', 'AI')} 초안이다.")
    _note(doc, "※ 제출: 재해 발생일부터 1개월 이내 (산업안전보건법 제57조제3항 · 시행규칙 제73조제1항). "
               "보존: 3년 (산업안전보건법 제164조제1항제4호, 사본 보존으로 시행규칙 제72조의 기록을 갈음).")
    return _save_docx(doc, path)
